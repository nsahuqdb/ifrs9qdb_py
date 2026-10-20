"""
One customer's ECL as an Excel calculator: the run's inputs, and every step
from them to the provision as a live formula.

``customer_ecl_workbook(run_dir, customer_id)`` returns the .xlsx. A reviewer
opens it, filters nothing, and can follow each figure of the report back to
the files the run priced:

    Lines       the customer's positions as the run read them (AccountMaster,
                CustomerStagingFlag), every value an input
    Overrides   the rating, stage and restructuring overrides applied to the
                customer in the run, with who, when and why
    Staging     the staging rules, the manual stage and customer contagion
    LGD         collateral value x allocation x (1 - haircut), then the LGD
    PD inputs   scenarios and weights, the macro model and its forecasts, the
                GCC history and forecast, and the scaling factor per scenario
                and year
    PD annual   per rating: point-in-time PD, cumulative and marginal PD by
                scenario, and the scenario-weighted marginal, years 1-50
    PD monthly  the monthly cumulative curve rebuilt from the annual chain,
                beside the run's own StPD.csv and the difference
    EAD         months to maturity, the amortisation shape and the horizon
    ECL         month by month: EAD x LGD x marginal PD / (1 + EIR)^(t/12)
    Summary     the provision per position and in total, against the report

THE BOUNDARY
    The workbook starts from the run's LIC input files (the Output folder),
    its frozen config and its override log -- the same point the report is
    priced from. How the ETL built those files from the raw extracts is the
    ETL's own record (validation and readiness reports), not repeated here.
    Two inputs are taken as given rather than derived: the scenario weights
    (the resolved values, with the mode that produced them) and a repayment
    schedule's monthly EAD (LifeTimeParameterOther).

ONE ANSWER
    The figures come from :func:`etl.report.compute_ecl`'s own inputs
    (:func:`ecl_frame`, the run's :class:`EclContext`) and the formulas
    follow it rule for rule, so the workbook's total equals the report's to
    the cent. The Summary says so per position; a difference is a finding.

    The PD curve the ECL tab reads is the rebuilt chain when it reproduces
    the run's StPD.csv (every run this code produced), else the run's StPD.
    The choice is a cell on the Settings tab and can be switched.
"""
from __future__ import annotations

import io
import math
import re
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

from .ids import as_id

__all__ = ["customer_ecl_facts", "write_customer_workbook",
           "customer_ecl_workbook", "run_customers"]

MAX_MONTH = 600          # EclConfig.max_month: the PD curves run to month 600
S1_MONTHS = 12           # Stage 1 looks twelve months ahead
MIN_MONTHS = 3           # months for a facility at or past maturity
PD_CAP = 1.0             # convert_to_monthly_stpd's cap
PD_SOURCES = ("Rebuilt chain", "Run StPD")


# ======================================================================
# Facts: everything the workbook shows, gathered with the engine's own
# functions so the numbers are the ones the report priced.
# ======================================================================
def _out_dir(run_dir: Path) -> Path:
    return run_dir / "Output" if (run_dir / "Output").is_dir() else run_dir


def _run_root(run_dir) -> Path:
    p = Path(run_dir)
    return p.parent if p.name == "Output" and not (p / "Output").is_dir() else p


def _str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and math.isnan(v):
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none", "<na>", "nat") else s


def _pay_type(v) -> str:
    """resolve_ead_shape()'s reading of a payment type."""
    s = _str(v)
    return s[:-2] if s.endswith(".0") else s


def _manual_stage(v):
    """classify_stage_report()'s reading of AccountMaster.Stage: its digits,
    when they are 1, 2 or 3."""
    s = re.sub(r"[^0-9]", "", _str(v))
    try:
        n = int(s)
    except ValueError:
        return None
    return n if n in (1, 2, 3) else None


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def _date(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    t = pd.Timestamp(v)
    return t.to_pydatetime().replace(tzinfo=None)


def run_customers(run_dir) -> pd.DataFrame:
    """Every customer in a run's books: id, name, positions, book(s) -- what
    the app's customer picker offers."""
    from .etl.report import read_output_csv, _chr, _col
    out_dir = _out_dir(_run_root(run_dir))
    rows = []
    for suffix, book in (("_1", "Lending"), ("_2", "Investments")):
        am = read_output_csv(out_dir, f"AccountMaster{suffix}.csv")
        if am is None or len(am) == 0:
            continue
        cm = read_output_csv(out_dir, f"CustomerMaster{suffix}.csv")
        cust = as_id(_chr(_col(am, "CustomerId")))
        names = {}
        if cm is not None and _col(cm, "CustomerName") is not None:
            k = as_id(_chr(_col(cm, "CustomerId")))
            for a, b in zip(k, _chr(_col(cm, "CustomerName"))):
                names.setdefault(a, _str(b))
        bal = pd.to_numeric(_col(am, "OnBalance"), errors="coerce")
        d = pd.DataFrame({"customer_id": cust, "on_balance": bal})
        d = d[d["customer_id"].notna() & (d["customer_id"] != "")]
        g = d.groupby("customer_id").agg(positions=("on_balance", "size"),
                                         on_balance=("on_balance", "sum"))
        g["book"] = book
        g["customer_name"] = [names.get(c, "") for c in g.index]
        rows.append(g.reset_index())
    if not rows:
        return pd.DataFrame(columns=["customer_id", "customer_name", "book",
                                     "positions", "on_balance"])
    d = pd.concat(rows, ignore_index=True)
    out = d.groupby("customer_id", as_index=False).agg(
        customer_name=("customer_name", lambda s: next((x for x in s if x), "")),
        book=("book", lambda s: " + ".join(dict.fromkeys(s))),
        positions=("positions", "sum"), on_balance=("on_balance", "sum"))
    return out.sort_values("on_balance", ascending=False).reset_index(drop=True)


def _report_rows(out_dir: Path, name: str) -> pd.DataFrame | None:
    p = out_dir / name
    if not p.is_file():
        return None
    return pd.read_csv(p, low_memory=False, dtype=str, keep_default_na=False)


def _overlay_report(out_dir: Path) -> tuple[pd.DataFrame | None, str | None]:
    """The overlaid report the app's analytics read (the first by name)."""
    ov = sorted(out_dir.glob("FinalEclReport_overlay_*.csv"))
    if not ov:
        return None, None
    return _report_rows(out_dir, ov[0].name), ov[0].name


def _overrides(run_root: Path, customer: str) -> pd.DataFrame:
    cols = ["kind", "value", "prior_value", "reason", "created_by",
            "created_at", "effect"]
    effect = {
        "rating": "Every contract of the customer takes this rating "
                  "(AccountMaster_1.Rating).",
        "stage": "Stage 3 sets IsDefault, Stage 2 sets IsLocal3 "
                 "(CustomerStagingFlag_1).",
        "restructuring": "Restructured sets IsLocal1 (CustomerStagingFlag_1).",
    }
    rows = []
    for kind in ("rating", "stage", "restructuring"):
        p = run_root / "overrides" / f"{kind}_overrides.csv"
        if not p.is_file():
            continue
        try:
            d = pd.read_csv(p, dtype=str, keep_default_na=False)
        except Exception:
            continue
        if "customer_id" not in d.columns:
            continue
        hit = d[as_id(d["customer_id"]) == customer]
        for _, r in hit.iterrows():
            rows.append({"kind": kind, "value": r.get("value", ""),
                         "prior_value": r.get("prior_value", ""),
                         "reason": r.get("reason", ""),
                         "created_by": r.get("created_by", ""),
                         "created_at": r.get("created_at", ""),
                         "effect": effect[kind]})
    return pd.DataFrame(rows, columns=cols)


def _collateral(out_dir: Path, cids: list[str]) -> dict:
    """The allocation rows of these contracts, and the collateral and type
    rows they point at (first occurrence, as the engine looks them up)."""
    from .etl.report import read_output_csv, _chr, _col, _num as rnum
    alloc = read_output_csv(out_dir, "AccountCollateralAllocation.csv")
    coll = read_output_csv(out_dir, "Collateral.csv")
    ctype = read_output_csv(out_dir, "CollateralType.csv")
    empty = {"alloc": pd.DataFrame(columns=["contract", "collateral", "share"]),
             "coll": pd.DataFrame(columns=["collateral", "type", "value"]),
             "types": pd.DataFrame(columns=["type", "haircut", "name"])}
    if alloc is None or coll is None or ctype is None:
        return empty
    a = pd.DataFrame({"contract": _chr(_col(alloc, "ContractId")),
                      "collateral": _chr(_col(alloc, "CollateralId")),
                      "share": rnum(_col(alloc, "AllocationPercentage"))})
    a = a[a["contract"].isin(set(cids))].reset_index(drop=True)
    c = pd.DataFrame({"collateral": _chr(_col(coll, "CollateralId")),
                      "type": _chr(_col(coll, "CollateralTypeId")),
                      "value": rnum(_col(coll, "CollateralValue"))})
    c = c[c["collateral"].notna()].drop_duplicates("collateral", keep="first")
    c = c[c["collateral"].isin(set(a["collateral"].dropna()))]
    name_col = next((n for n in ("CollateralTypeName", "Name", "Description")
                     if _col(ctype, n) is not None), None)
    t = pd.DataFrame({"type": _chr(_col(ctype, "CollateralTypeId")),
                      "haircut": rnum(_col(ctype, "HaircutGeneral")),
                      "name": (_chr(_col(ctype, name_col)) if name_col
                               else pd.Series([None] * len(ctype)))})
    t = t[t["type"].notna()].drop_duplicates("type", keep="first")
    t = t[t["type"].isin(set(c["type"].dropna()))]
    return {"alloc": a, "coll": c.reset_index(drop=True),
            "types": t.reset_index(drop=True)}


def _pd_chain_inputs(run_root: Path, static, model_cfg) -> dict | None:
    """What build_stpd_from_static() builds the run's StPD from, resolved
    exactly as it resolves it. None when the run's config cannot rebuild it."""
    import yaml
    from .etl import macro as M
    from .etl.model_registry import mev_model_weights, resolve_model, run_model_id
    p = run_root / "config_used" / "config" / "model_inputs.yml"
    if not p.is_file() or static is None or model_cfg is None:
        return None
    model_inputs = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    resolved = resolve_model(model_cfg, run_model_id(run_root))
    mevs = resolved["model"]["mevs"]
    weights = mev_model_weights(resolved, model_inputs)
    fc_block = model_inputs["mev_forecasts"]["forecasts"]
    years = sorted(fc_block, key=lambda k: int(k))
    forecasts = np.array([np.asarray(fc_block[y], dtype=float) for y in years])
    mc = model_cfg
    horizons = mc.get("horizons", {}) or {}
    n_fc = int(horizons.get("n_forecasts", mc.get("n_forecasts", forecasts.shape[0])))
    max_mat = int(horizons.get("max_maturity", mc.get("max_maturity", 50)))
    scen = static["scenario_severity"]
    gcc_hist, ext_fc = M._gcc_inputs(static, model_inputs, n_fc)
    w_int, w_ext = M.resolve_scenario_weights(static, model_cfg, model_inputs)
    ext_on = bool(gcc_hist.size) and ext_fc is not None
    out = {
        "model_id": resolved["model"].get("id"),
        "anchor": float(mc["ttc_anchor_pd"]),
        "n_forecasts": n_fc, "max_maturity": max_mat,
        "mevs": [{"name": m.get("name") or m.get("variable_id"),
                  "intercept": float(m["intercept"]),
                  "coefficient": float(m["coefficient"]),
                  "sd": float(m["standard_deviation"]),
                  "unit": float(m["stress_unit_multiplier"]),
                  "weight": float(w)} for m, w in zip(mevs, weights)],
        "forecast_years": [int(y) for y in years],
        "forecasts": forecasts,
        "scenarios": [(str(r.scenario), float(r.severity_z))
                      for r in scen.itertuples()],
        "w_int": {str(k): float(v) for k, v in w_int.items()},
        "w_int_mode": (model_inputs.get("internal_scenario_weights") or {}).get(
            "mode", "explicit"),
        "w_ext": w_ext, "w_ext_mode": (model_inputs.get(
            "external_scenario_weights") or {}).get("mode", "explicit"),
        "ext_on": ext_on,
        "gcc_hist": np.asarray(gcc_hist, dtype=float),
        "ext_fc": None if ext_fc is None else np.asarray(ext_fc, dtype=float),
        "ttc": {1: static.ttc_for(1), 2: static.ttc_for(2)},
        "ratings": {1: static.ratings_for(1), 2: static.ratings_for(2)},
        "gcc_tables": _gcc_tables(static, model_inputs, ext_fc),
    }
    return out


def _gcc_tables(static, model_inputs, ext_fc) -> dict:
    """The per-country tables behind the GCC history and forecast, laid out
    as gcc_weighted_history() and external_gcc_forecast() read them."""
    def col(df, *names):
        low = {c.lower(): c for c in df.columns}
        for n in names:
            if n in low:
                return df[low[n]]
        return None
    out = {"hist": None, "fc": None}
    g, p = static.get("gcc_real_gdp_growth"), static.get("gcc_gdp_current_prices")
    if g is not None and p is not None and len(g) and len(p) \
            and col(g, "country") is not None and col(p, "country") is not None:
        from .etl.macro import GCC_HISTORY_FIRST_YEAR
        gg = pd.DataFrame({"year": pd.to_numeric(col(g, "year"), errors="coerce"),
                           "country": col(g, "country"),
                           "growth": pd.to_numeric(col(g, "value", "growth"),
                                                   errors="coerce")}).dropna()
        pp = pd.DataFrame({"year": pd.to_numeric(col(p, "year"), errors="coerce"),
                           "country": col(p, "country"),
                           "gdp": pd.to_numeric(col(p, "value", "gdp"),
                                                errors="coerce")}).dropna()
        gg = gg[gg["year"] >= GCC_HISTORY_FIRST_YEAR]
        m = gg.merge(pp, on=["year", "country"], how="inner")
        n = m.groupby("year").size()
        m = m[m["year"].isin(n[n >= 2].index)]
        if len(m):
            countries = list(dict.fromkeys(m.sort_values("country")["country"]))
            yrs = sorted(m["year"].unique())
            gw = m.pivot(index="year", columns="country", values="growth").reindex(
                index=yrs, columns=countries)
            pw = m.pivot(index="year", columns="country", values="gdp").reindex(
                index=yrs, columns=countries)
            out["hist"] = {"years": [int(y) for y in yrs], "countries": countries,
                           "growth": gw.to_numpy(), "gdp": pw.to_numpy()}
    block = model_inputs.get("external_gcc_country_growth")
    prices = model_inputs.get("external_gcc_country_prices")
    if (model_inputs.get("external_gcc_forecast") is None and block and prices
            and ext_fc is not None):
        countries = list(block)
        n = len(np.atleast_1d(np.asarray(ext_fc, dtype=float)))
        gr = np.full((n, len(countries)), np.nan)
        pr = np.full((n, len(countries)), np.nan)
        for j, c in enumerate(countries):
            v = np.atleast_1d(np.asarray(block[c], dtype=float))[:n]
            gr[:v.size, j] = v
            w = np.atleast_1d(np.asarray(prices.get(c, []), dtype=float))[:n]
            pr[:w.size, j] = w
        out["fc"] = {"countries": countries, "growth": gr, "gdp": pr}
    return out


def _chain_rating(chain: dict, rt: int, bucket) -> tuple[str, float] | None:
    """The TTC grade whose curve StPD carries for this bucket: the grade with
    that hierarchy, the last of them if several (build_stpd writes the grades
    in name order and the curve lookup keeps the last)."""
    if chain is None or bucket is None:
        return None
    ttc = chain["ttc"][rt]
    rat = chain["ratings"][rt]
    h = dict(zip(rat["rating"].astype(str),
                 pd.to_numeric(rat["hierarchy"], errors="coerce")))
    try:
        b = float(bucket)
    except (TypeError, ValueError):
        return None
    hits = sorted(str(r) for r in ttc["rating"] if h.get(str(r)) == b)
    if not hits:
        return None
    r = hits[-1]
    v = pd.to_numeric(ttc.loc[ttc["rating"].astype(str) == r, "ttc_pd"],
                      errors="coerce").iloc[0]
    return r, float(v)


def customer_ecl_facts(run_dir, customer_id) -> dict:
    """Everything the calculator shows for one customer of a run."""
    from . import __version__
    from .etl.report import (build_ecl_context, compute_ecl, ecl_frame,
                             read_output_csv, run_pricing_config, _chr, _col,
                             _portfolio_map, _KEYSEP)
    from .stress import staging_threshold
    from .engine import ead_fallback_rules

    root = _run_root(run_dir)
    out_dir = _out_dir(root)
    cust = as_id(pd.Series([customer_id])).iloc[0]
    static, model_cfg = run_pricing_config(root)
    pmap = _portfolio_map(static)
    thr = float(staging_threshold(out_dir))
    ctx = build_ecl_context(out_dir, model_cfg=model_cfg)
    frame = ecl_frame(out_dir, pmap, thr)
    if frame is None:
        raise FileNotFoundError(f"No AccountMaster output in {out_dir}")
    sel = frame[as_id(frame["customer"]) == cust]
    if len(sel) == 0:
        raise KeyError(f"Customer {customer_id!r} has no position in this run")
    priced = compute_ecl(ctx, sel)

    # the run's report, row for row: the report and the frame keep the
    # AccountMaster order, lending then investments
    rep = _report_rows(out_dir, "FinalEclReport.csv")
    ovr, ovr_name = _overlay_report(out_dir)

    def rep_get(df, i, col):
        if df is None or col not in df.columns or i >= len(df):
            return None
        return df.iloc[i][col]

    # raw columns the frame does not keep
    am = {s: read_output_csv(out_dir, f"AccountMaster{s}.csv") for s in ("_1", "_2")}
    cm = {s: read_output_csv(out_dir, f"CustomerMaster{s}.csv") for s in ("_1", "_2")}
    n1 = 0 if am["_1"] is None else len(am["_1"])

    lines = []
    for k, (i, r) in enumerate(sel.iterrows(), start=1):
        s = r["suffix"]
        j = i if s == "_1" else i - n1
        a = am[s]
        stage_raw = _str(_chr(_col(a, "Stage")).iloc[j]) if _col(a, "Stage") is not None else ""
        if rep is not None and i < len(rep):
            rc = as_id(pd.Series([rep.iloc[i]["Contract Id"]])).iloc[0]
            if rc != r["contract"]:
                rep_i = None
            else:
                rep_i = i
        else:
            rep_i = None
        p = priced.loc[i]
        lines.append({
            "line": k, "book": r["book"], "suffix": s, "cid": _str(r["cid"]),
            "contract": _str(r["contract"]), "customer": cust,
            "account_type": _str(r["account_type"]),
            "portfolio": _str(r["portfolio"]),
            "rating_type": _str(r["rating_type"]), "rating": _str(r["rating"]),
            "d_extract": _date(r["d_extract"]), "d_maturity": _date(r["d_maturity"]),
            "on_balance": _num(r["on_balance"]), "off_balance": _num(r["off_balance"]),
            "eir": _num(r["eir"]), "dpd": _num(r["dpd"]),
            "payment_type": _pay_type(r["payment_type"]),
            "pay_freq": _num(r["pay_freq"]), "deferral": _num(r["deferral"]),
            "nir": _num(r["nir"]), "stage_raw": stage_raw,
            "stage_manual": _manual_stage(stage_raw),
            **{f: _num(r[f]) for f in ("IsWatchlist", "IsDefault", "IsDefaultInGCC",
                                       "IsInsolvency", "IsLocal1", "IsLocal2",
                                       "IsLocal3", "IsLocal4", "IsLocal5",
                                       "IsLocal6")},
            "schedule": ctx.schedules.get(r["cid"]),
            "stage_engine": int(r["stage"]),
            "ecl_engine": _num(p["ecl"]),
            "rep_stage": _num(rep_get(rep, rep_i, "Ifrs Stage")) if rep_i is not None else None,
            # the model figure: the overlay waterfall's first column when the
            # report has one, else the provision column
            "rep_ecl": _num(rep_get(rep, rep_i, "Ecl Model Onbal")
                            if rep is not None and "Ecl Model Onbal" in rep.columns
                            and _str(rep_get(rep, rep_i, "Ecl Model Onbal")) != ""
                            else rep_get(rep, rep_i, "Cla Amount Onbal"))
            if rep_i is not None else None,
            "rep_pdlife": _num(rep_get(rep, rep_i, "PD Lifetime Value")) if rep_i is not None else None,
            "rep_lgd": _num(rep_get(rep, rep_i, "LGD Rate")) if rep_i is not None else None,
            # the overlay: the overlaid report the analytics read when there
            # is one, else the report's own waterfall
            "overlay_id": _str(rep_get(ovr if ovr is not None else rep, rep_i,
                                       "Overlay Id")) if rep_i is not None else "",
            "overlay": _num(rep_get(ovr if ovr is not None else rep, rep_i,
                                    "Overlay Amount")) if rep_i is not None else None,
            "rep_final": _num(rep_get(ovr if ovr is not None else rep, rep_i,
                                      "Cla Amount Onbal")) if rep_i is not None else None,
        })

    # customer name, from the book(s) the customer is in
    name = ""
    for s in ("_1", "_2"):
        c = cm[s]
        if c is None or _col(c, "CustomerName") is None:
            continue
        hit = c[as_id(_chr(_col(c, "CustomerId"))) == cust]
        if len(hit):
            name = _str(_chr(_col(hit, "CustomerName")).iloc[0])
            break

    # ratings -> buckets, as the curve lookup reads Ratings.csv
    rat = read_output_csv(out_dir, "Ratings.csv")
    rmap = []
    if rat is not None:
        rd = pd.DataFrame({"rt": _chr(_col(rat, "RatingType")).astype(str),
                           "rating": _chr(_col(rat, "Rating")).astype(str),
                           "bucket": _chr(_col(rat, "Hierarchy"))})
        want = {("2" if (l["rating_type"].startswith("External")
                         or l["rating_type"] == "2") else "1", l["rating"])
                for l in lines}
        seen = set()
        for row in rd.itertuples():
            key = (row.rt, row.rating)
            if key in want and key not in seen:
                seen.add(key)
                rmap.append({"rt": row.rt, "rating": row.rating,
                             "bucket": _num(row.bucket)})

    # the PD chain and the curves each line prices on
    try:
        chain = _pd_chain_inputs(root, static, model_cfg)
    except Exception as exc:              # a run whose config cannot rebuild
        chain = None
        chain_error = f"{type(exc).__name__}: {exc}"
    else:
        chain_error = None if chain is not None else \
            "the run froze no model_inputs.yml"
    bucket_of = {(m["rt"], m["rating"]): m["bucket"] for m in rmap}
    curves = []
    for l in lines:
        rt = "2" if (l["rating_type"].startswith("External")
                     or l["rating_type"] == "2") else "1"
        b = bucket_of.get((rt, l["rating"]))
        bkey = None if b is None else (str(int(b)) if float(b).is_integer() else str(b))
        l["rt"], l["bucket"] = rt, b
        key = (l["portfolio"], bkey)
        l["curve_key"] = f"{l['portfolio']}|{bkey}" if bkey is not None else None
        if bkey is None or any(c["key"] == key for c in curves):
            continue
        stpd = ctx.stpd.get(f"{l['portfolio']}{_KEYSEP}{bkey}")
        curves.append({"key": key, "label": f"{l['portfolio']}|{bkey}",
                       "portfolio": l["portfolio"], "bucket": b, "rt": int(rt),
                       "stpd": None if stpd is None else np.asarray(stpd, dtype=float),
                       "chain": _chain_rating(chain, int(rt), b)})

    rebuilt = _rebuild_curves(chain, curves) if chain is not None else {}
    worst = 0.0
    for c in curves:
        rb, sp = rebuilt.get(c["key"]), c["stpd"]
        if rb is None or sp is None:
            worst = float("inf")
            continue
        n = min(len(rb), len(sp))
        worst = max(worst, float(np.nanmax(np.abs(rb[:n] - sp[:n]))))
    pd_source = PD_SOURCES[0] if (curves and worst < 1e-9) else PD_SOURCES[1]

    rules, default = ead_fallback_rules(model_cfg)
    return {
        "run_id": root.name, "run_dir": str(root),
        "customer_id": cust, "customer_name": name,
        "engine": f"ifrs9qdb (python) {__version__}",
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "report_file": "FinalEclReport.csv" if rep is not None else None,
        "overlay_file": ovr_name,
        "settings": {"threshold": thr, "lgd_base": ctx.lgd_base,
                     "lgd_floor": ctx.lgd_floor, "zero_lgd": ctx.cfg.zero_exposure_lgd,
                     "stage3": ctx.stage3, "cap": bool(ctx.cap),
                     "rules": [(None if pf is None else str(pf), str(pt), str(sh))
                               for pf, pt, sh in rules],
                     "default_shape": str(default)},
        "lines": lines, "overrides": _overrides(root, cust),
        "collateral": _collateral(out_dir, [l["cid"] for l in lines]),
        "ratings_map": rmap, "chain": chain, "chain_error": chain_error,
        "curves": curves, "rebuilt": rebuilt, "pd_source": pd_source,
        "rebuild_gap": worst,
    }


def _rebuild_curves(chain: dict, curves: list) -> dict:
    """Each curve rebuilt with the package's own chain, to decide whether the
    workbook's rebuilt curve can stand in for the run's StPD."""
    from .etl import macro as M
    out = {}
    scen = pd.DataFrame(chain["scenarios"], columns=["scenario", "severity_z"])
    specs = [{"standard_deviation": m["sd"], "stress_unit_multiplier": m["unit"],
              "intercept": m["intercept"], "coefficient": m["coefficient"]}
             for m in chain["mevs"]]
    w = [m["weight"] for m in chain["mevs"]]
    for c in curves:
        if c["chain"] is None:
            continue
        r, ttc = c["chain"]
        t = pd.DataFrame({"rating": [r], "ttc_pd": [ttc]})
        if c["rt"] == 2 and chain["ext_on"]:
            ts = M.build_term_structure_external(t, scen, chain["ext_fc"],
                                                 chain["gcc_hist"],
                                                 chain["max_maturity"])
        else:
            ts = M.build_term_structure(t, scen, specs, w, chain["forecasts"],
                                        chain["anchor"], chain["n_forecasts"],
                                        chain["max_maturity"])
        wts = chain["w_ext"] if c["rt"] == 2 else chain["w_int"]
        if isinstance(wts, dict) and "per_year" in wts:
            ann = M.apply_scenario_weights_per_year(ts, wts["per_year"], wts["average"])
        else:
            ann = M.apply_scenario_weights(ts, wts)
        mo = M.convert_to_monthly_stpd(ann, MAX_MONTH)
        out[c["key"]] = np.concatenate([[0.0], mo["pd_lifetime"].to_numpy()])
    return out


# ======================================================================
# The workbook
# ======================================================================
def _col_letter(c: int) -> str:
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def _a(r: int, c: int, sheet: str | None = None, fixed: bool = True) -> str:
    """A cell address, absolute by default, on another sheet when named."""
    ref = (f"${_col_letter(c)}${r}" if fixed else f"{_col_letter(c)}{r}")
    return f"'{sheet}'!{ref}" if sheet else ref


def _rng(r1, c1, r2, c2, sheet=None) -> str:
    ref = f"${_col_letter(c1)}${r1}:${_col_letter(c2)}${r2}"
    return f"'{sheet}'!{ref}" if sheet else ref


class _Book:
    """A thin layer over openpyxl: inputs blue, formulas black, one header
    style, number formats by kind."""

    FMT = {"amt": "#,##0.00", "pct": "0.0000%", "pd": "0.000000000",
           "rate": "0.000000", "int": "0", "date": "yyyy-mm-dd", "num": "General",
           "x": "0.0000"}

    def __init__(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.Font, self.Fill, self.Align = Font, PatternFill, Alignment
        self.f_in = Font(color="0000FF")
        self.f_hdr = Font(bold=True, color="FFFFFF")
        self.fill_hdr = PatternFill("solid", fgColor="305496")
        self.f_title = Font(bold=True, size=14)
        self.f_sub = Font(bold=True, size=11)
        self.f_note = Font(italic=True, color="595959")
        self.fill_key = PatternFill("solid", fgColor="FFF2CC")

    def sheet(self, name):
        return self.wb.create_sheet(name)

    def put(self, ws, r, c, v, kind=None, inp=False, bold=False, key=False):
        cell = ws.cell(r, c)
        if isinstance(v, float) and not math.isfinite(v):
            v = None
        cell.value = v
        if kind:
            cell.number_format = self.FMT[kind]
        if inp and not (isinstance(v, str) and v.startswith("=")):
            cell.font = self.f_in
        if bold:
            cell.font = self.Font(bold=True)
        if key:
            cell.fill = self.fill_key
            cell.font = self.Font(bold=True)
        return cell

    def header(self, ws, r, c, labels):
        for k, lab in enumerate(labels):
            cell = ws.cell(r, c + k, lab)
            cell.font, cell.fill = self.f_hdr, self.fill_hdr
            cell.alignment = self.Align(wrap_text=True, vertical="top")

    def title(self, ws, text, note=None):
        ws.cell(1, 1, text).font = self.f_title
        if note:
            ws.cell(2, 1, note).font = self.f_note

    def note(self, ws, r, c, text):
        ws.cell(r, c, text).font = self.f_note

    def widths(self, ws, w: dict):
        for col, width in w.items():
            ws.column_dimensions[col].width = width


def write_customer_workbook(facts: dict, dest=None):
    """Write the calculator for ``facts`` (:func:`customer_ecl_facts`) to
    ``dest`` (a path or a binary buffer); returns the bytes when ``dest`` is
    None."""
    B = _Book()
    L = facts["lines"]
    st = facts["settings"]
    chain = facts["chain"]
    curves = facts["curves"]
    ext_chain = bool(chain and chain["ext_on"])

    ws_sum = B.sheet("Summary")
    ws_set = B.sheet("Settings")
    ws_ln = B.sheet("Lines")
    ws_ov = B.sheet("Overrides")
    ws_stg = B.sheet("Staging")
    ws_lgd = B.sheet("LGD")
    ws_pin = B.sheet("PD inputs")
    ws_pa = B.sheet("PD annual")
    ws_pm = B.sheet("PD monthly")
    ws_ead = B.sheet("EAD")
    ws_ecl = B.sheet("ECL")

    # ---------------------------------------------------------- Settings
    B.title(ws_set, "Settings the run priced with",
            "From the run's frozen config (config_used). Blue cells are inputs; "
            "change one and every formula follows.")
    S = {}
    rows = [
        ("thr", "Stage 2 DPD threshold (days, DPD > threshold)", st["threshold"],
         "int", "staging_thresholds.csv; 60 when the run froze none"),
        ("s3dpd", "Stage 3 DPD (days, DPD > this)", 90, "int",
         "fixed in the staging rule"),
        ("lgd_base", "LGD base", st["lgd_base"], "x", "model.yml ecl.lgd_base"),
        ("lgd_floor", "Unsecured floor", st["lgd_floor"], "x",
         "model.yml ecl.lgd_unsecured_floor: at least this share is unsecured"),
        ("zero_lgd", "LGD at zero or missing exposure", st["zero_lgd"], "x", ""),
        ("stage3", "Stage 3 method", st["stage3"], None,
         "full_outstanding: Stage 3 ECL = on-balance exposure; zero: nil"),
        ("cap", "Cap ECL at on-balance exposure", bool(st["cap"]), None,
         "model.yml ecl.cap_ecl_at_exposure"),
        ("s1m", "Stage 1 horizon (months)", S1_MONTHS, "int", "12-month ECL"),
        ("minm", "Months for a facility at or past maturity", MIN_MONTHS, "int",
         "LIC's floor for a run-off facility"),
        ("maxm", "Last month of the PD curves", MAX_MONTH, "int", ""),
        ("pdcap", "Cap on the cumulative PD", PD_CAP, "x", ""),
        ("pdsrc", "PD curve the ECL uses", facts["pd_source"], None,
         "'Rebuilt chain' (PD annual -> PD monthly) or 'Run StPD' (the run's "
         "StPD.csv). Rebuilt when it reproduces StPD.csv."),
    ]
    B.header(ws_set, 4, 1, ["Setting", "Value", "Note"])
    r = 5
    for key, lab, val, kind, note in rows:
        B.put(ws_set, r, 1, lab)
        B.put(ws_set, r, 2, val, kind, inp=True)
        B.note(ws_set, r, 3, note)
        S[key] = _a(r, 2, "Settings")
        r += 1
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Rebuilt chain,Run StPD"',
                        allow_blank=False)
    ws_set.add_data_validation(dv)
    dv.add(ws_set.cell(r - 1, 2))
    r += 1
    ws_set.cell(r, 1, "EAD fallback: amortisation shape for a position with "
                      "no repayment schedule").font = B.f_sub
    B.note(ws_set, r + 1, 1, "First a rule naming the position's portfolio "
           "and payment type, then a rule with no portfolio, else the default. "
           "A facility with 3 months or less to run is a bullet.")
    B.header(ws_set, r + 2, 1, ["Portfolio", "Payment type", "Shape", "Key"])
    r0 = r + 3
    for k, (pf, pt, sh) in enumerate(st["rules"]):
        rr = r0 + k
        B.put(ws_set, rr, 1, pf or "", inp=True)
        B.put(ws_set, rr, 2, pt, inp=True)
        B.put(ws_set, rr, 3, sh, inp=True)
        B.put(ws_set, rr, 4, f"={_a(rr, 1, fixed=False)}&\"|\"&{_a(rr, 2, fixed=False)}")
    r1 = r0 + max(len(st["rules"]), 1) - 1
    S["rule_key"] = _rng(r0, 4, r1, 4, "Settings")
    S["rule_shape"] = _rng(r0, 3, r1, 3, "Settings")
    B.put(ws_set, r1 + 2, 1, "Default shape")
    B.put(ws_set, r1 + 2, 2, st["default_shape"], inp=True)
    S["default_shape"] = _a(r1 + 2, 2, "Settings")
    B.widths(ws_set, {"A": 48, "B": 18, "C": 90, "D": 22})

    # ---------------------------------------------------------- Lines
    B.title(ws_ln, "The customer's positions, as the run read them",
            "AccountMaster_1/_2 and CustomerStagingFlag_1/_2 from the run's "
            "Output folder; a flag is the customer's, repeated on each of its "
            "positions in that book. Every value here is an input.")
    ln_cols = ["Line", "Book", "Contract", "Customer", "Account type",
               "Portfolio", "Rating type", "Rating", "Extract date",
               "Maturity date", "On balance", "Off balance", "EIR",
               "Past due days", "Payment type", "Payment frequency (months)",
               "Deferral (months)", "Nominal interest rate",
               "AccountMaster.Stage", "Manual stage (its digits)",
               "IsWatchlist", "IsDefault", "IsDefaultInGCC", "IsInsolvency",
               "IsLocal1", "IsLocal2", "IsLocal3", "IsLocal4", "IsLocal5",
               "IsLocal6", "Repayment schedule", "Rating bucket", "Curve no."]
    C = {n: i + 1 for i, n in enumerate(ln_cols)}
    HR = 4
    B.header(ws_ln, HR, 1, ln_cols)
    LR = {}                                        # line -> row on Lines
    for k, l in enumerate(L):
        rr = HR + 1 + k
        LR[l["line"]] = rr
        vals = [l["line"], l["book"], l["cid"], l["customer"], l["account_type"],
                l["portfolio"], l["rating_type"], l["rating"], l["d_extract"],
                l["d_maturity"], l["on_balance"], l["off_balance"], l["eir"],
                l["dpd"], l["payment_type"], l["pay_freq"], l["deferral"],
                l["nir"], l["stage_raw"], l["stage_manual"],
                l["IsWatchlist"], l["IsDefault"], l["IsDefaultInGCC"],
                l["IsInsolvency"], l["IsLocal1"], l["IsLocal2"], l["IsLocal3"],
                l["IsLocal4"], l["IsLocal5"], l["IsLocal6"],
                "Yes" if l["schedule"] is not None else "No"]
        kinds = [None, None, None, None, None, None, None, None, "date", "date",
                 "amt", "amt", "rate", "int", None, "int", "int", "rate", None,
                 "int"] + ["int"] * 10 + [None]
        for c, (v, kd) in enumerate(zip(vals, kinds), start=1):
            B.put(ws_ln, rr, c, v, kd, inp=c > 1)
    S["ln_first"], S["ln_last"] = HR + 1, HR + len(L)

    # ---------------------------------------------------------- PD inputs
    pin = _write_pd_inputs(B, ws_pin, facts, S)
    # rating bucket and curve number on Lines, by formula
    for l in L:
        rr = LR[l["line"]]
        rt = (f"IF(OR(LEFT({_a(rr, C['Rating type'], fixed=False)},8)=\"External\","
              f"{_a(rr, C['Rating type'], fixed=False)}=\"2\"),\"2\",\"1\")")
        B.put(ws_ln, rr, C["Rating bucket"],
              f"=IFERROR(INDEX({pin['rm_bucket']},MATCH({rt}&\"|\"&"
              f"{_a(rr, C['Rating'], fixed=False)},{pin['rm_key']},0)),\"\")", "int")
    # ---------------------------------------------------------- PD annual / monthly
    pa = _write_pd_annual(B, ws_pa, facts, pin)
    pm = _write_pd_monthly(B, ws_pm, facts, pa, S)
    for l in L:
        rr = LR[l["line"]]
        B.put(ws_ln, rr, C["Curve no."],
              f"=IFERROR(MATCH({_a(rr, C['Portfolio'], fixed=False)}&\"|\"&"
              f"{_a(rr, C['Rating bucket'], fixed=False)},{pm['keys']},0),\"\")",
              "int")
    B.widths(ws_ln, {"A": 6, "B": 11, "C": 14, "D": 12, "E": 12, "F": 16,
                     "G": 11, "H": 10, "I": 12, "J": 12, "K": 16, "L": 14,
                     "M": 11, "N": 9})
    ws_ln.freeze_panes = ws_ln.cell(HR + 1, 4)

    def ln(line, name, fixed=True):
        return _a(LR[line], C[name], "Lines", fixed)

    # ---------------------------------------------------------- Overrides
    B.title(ws_ov, "Overrides applied to this customer in the run",
            "From the run's overrides folder (reviewed at the pause step). They "
            "are already in the Lines tab's rating and flags: this is the record "
            "of what changed and why.")
    B.header(ws_ov, 4, 1, ["Kind", "New value", "Model value (before)",
                           "Reason", "By", "When", "Where it shows"])
    ov = facts["overrides"]
    if len(ov) == 0:
        B.note(ws_ov, 5, 1, "No rating, stage or restructuring override for this "
                            "customer in this run.")
    for k, row in enumerate(ov.itertuples(index=False)):
        for c, v in enumerate([row.kind, row.value, row.prior_value, row.reason,
                               row.created_by, row.created_at, row.effect], 1):
            B.put(ws_ov, 5 + k, c, v, inp=True)
    rr = 7 + max(len(ov), 1)
    ws_ov.cell(rr, 1, "Manual stage in AccountMaster.Stage").font = B.f_sub
    B.note(ws_ov, rr + 1, 1, "A stage written into the account master by the "
           "extract. It replaces the rule stage unless the rules say Stage 3; "
           "see the Staging tab.")
    man = [l for l in L if l["stage_manual"] is not None]
    B.header(ws_ov, rr + 2, 1, ["Line", "Contract", "AccountMaster.Stage"])
    if not man:
        B.note(ws_ov, rr + 3, 1, "None on this customer's positions.")
    for k, l in enumerate(man):
        B.put(ws_ov, rr + 3 + k, 1, f"={ln(l['line'], 'Line')}")
        B.put(ws_ov, rr + 3 + k, 2, f"={ln(l['line'], 'Contract')}")
        B.put(ws_ov, rr + 3 + k, 3, f"={ln(l['line'], 'AccountMaster.Stage')}")
    B.widths(ws_ov, {"A": 15, "B": 14, "C": 20, "D": 60, "E": 12, "F": 26,
                     "G": 60})

    # ---------------------------------------------------------- Staging
    B.title(ws_stg, "Stage, rule by rule",
            "classify_stage_report(): DPD over the threshold (and not over 90), "
            "the watchlist or any local flag -> 2; Tasdeer -> 2; DPD over 90 or "
            "the default flag -> 3; a manual stage replaces the rule stage "
            "unless that is 3; then a customer with any Stage 2 or 3 position "
            "in the book lifts its Stage 1 positions to 2 (Tasdeer excepted).")
    sg_cols = ["Line", "Contract", "Book", "Portfolio", "DPD", "DPD rule",
               "Watchlist", "Any local flag (or watchlist)", "Tasdeer",
               "Stage 3 rule", "Rule stage", "Manual stage", "After manual",
               "Customer's worst in book", "Final stage", "Why",
               "Report stage", "Check"]
    G = {n: i + 1 for i, n in enumerate(sg_cols)}
    B.header(ws_stg, 4, 1, sg_cols)
    GR = {}
    for k, l in enumerate(L):
        GR[l["line"]] = 5 + k
    books = {}
    for l in L:
        books.setdefault(l["suffix"], []).append(GR[l["line"]])
    for l in L:
        rr, x = GR[l["line"]], l["line"]
        g = lambda name: _a(rr, G[name], fixed=False)
        dpd = ln(x, "Past due days")
        B.put(ws_stg, rr, G["Line"], f"={ln(x, 'Line')}")
        B.put(ws_stg, rr, G["Contract"], f"={ln(x, 'Contract')}")
        B.put(ws_stg, rr, G["Book"], f"={ln(x, 'Book')}")
        B.put(ws_stg, rr, G["Portfolio"], f"={ln(x, 'Portfolio')}")
        B.put(ws_stg, rr, G["DPD"], f"=IF(ISNUMBER({dpd}),{dpd},\"\")", "int")
        B.put(ws_stg, rr, G["DPD rule"],
              f"=AND(ISNUMBER({dpd}),N({dpd})>{S['thr']},N({dpd})<={S['s3dpd']})")
        B.put(ws_stg, rr, G["Watchlist"], f"={ln(x, 'IsWatchlist')}=1")
        flags = ",".join(f"{ln(x, f)}=1" for f in
                         ("IsWatchlist", "IsLocal1", "IsLocal2", "IsLocal3",
                          "IsLocal4", "IsLocal5", "IsLocal6"))
        B.put(ws_stg, rr, G["Any local flag (or watchlist)"], f"=OR({flags})")
        B.put(ws_stg, rr, G["Tasdeer"], f"={ln(x, 'Portfolio')}=\"Tasdeer\"")
        B.put(ws_stg, rr, G["Stage 3 rule"],
              f"=OR(AND(ISNUMBER({dpd}),N({dpd})>{S['s3dpd']}),{ln(x, 'IsDefault')}=1)")
        B.put(ws_stg, rr, G["Rule stage"],
              f"=IF({g('Stage 3 rule')},3,IF(OR({g('DPD rule')},{g('Watchlist')},"
              f"{g('Any local flag (or watchlist)')},{g('Tasdeer')}),2,1))", "int")
        man = ln(x, "Manual stage (its digits)")
        B.put(ws_stg, rr, G["Manual stage"],
              f"=IF(ISNUMBER({man}),{man},\"\")", "int")
        B.put(ws_stg, rr, G["After manual"],
              f"=IF({g('Stage 3 rule')},3,IF(ISNUMBER({g('Manual stage')}),"
              f"{g('Manual stage')},{g('Rule stage')}))", "int")
        rows_b = books[l["suffix"]]
        B.put(ws_stg, rr, G["Customer's worst in book"],
              f"=MAX({_rng(min(rows_b), G['After manual'], max(rows_b), G['After manual'])})",
              "int")
        worst = _a(rr, G["Customer's worst in book"], fixed=False)
        B.put(ws_stg, rr, G["Final stage"],
              f"=IF(AND({g('After manual')}=1,{worst}>=2,"
              f"NOT({g('Tasdeer')})),2,{g('After manual')})", "int", key=True)
        why = (f"=IF({g('Stage 3 rule')},IF({ln(x, 'IsDefault')}=1,\"Default flag\","
               f"\"DPD over \"&{S['s3dpd']}),IF(ISNUMBER({g('Manual stage')}),"
               f"\"Manual stage \"&{g('Manual stage')},IF({g('DPD rule')},"
               f"\"DPD over \"&{S['thr']},IF({g('Watchlist')},\"Watchlist\","
               f"IF({g('Any local flag (or watchlist)')},\"Local flag\","
               f"IF({g('Tasdeer')},\"Tasdeer (collective)\","
               f"IF({g('Final stage')}=2,\"Contagion: the customer has a Stage \"&"
               f"{worst}&\" position\","
               f"\"Performing\")))))))")
        B.put(ws_stg, rr, G["Why"], why)
        B.put(ws_stg, rr, G["Report stage"], l["rep_stage"], "int", inp=True)
        B.put(ws_stg, rr, G["Check"],
              f"=IF({g('Report stage')}=\"\",\"\",IF({g('Final stage')}="
              f"{g('Report stage')},\"OK\",\"DIFFERS\"))")
    B.widths(ws_stg, {"A": 6, "B": 14, "C": 11, "D": 16, "E": 7, "F": 9,
                      "G": 10, "H": 12, "I": 9, "J": 10, "K": 9, "L": 9,
                      "M": 9, "N": 11, "O": 9, "P": 42, "Q": 9, "R": 9})
    ws_stg.freeze_panes = ws_stg.cell(5, 3)

    def stage(line):
        return _a(GR[line], G["Final stage"], "Staging")

    # ---------------------------------------------------------- LGD
    lg = _write_lgd(B, ws_lgd, facts, L, ln, S)

    # ---------------------------------------------------------- EAD
    ead = _write_ead(B, ws_ead, facts, L, ln, stage, S)

    # ---------------------------------------------------------- ECL
    ecl = _write_ecl(B, ws_ecl, facts, L, ln, lg, ead, pm, S)

    # ---------------------------------------------------------- Summary
    _write_summary(B, ws_sum, facts, L, ln, stage, lg, ead, ecl, pm, S, GR, G)

    buf = io.BytesIO()
    B.wb.calculation.fullCalcOnLoad = True
    B.wb.save(buf)
    data = buf.getvalue()
    if dest is None:
        return data
    if hasattr(dest, "write"):
        dest.write(data)
        return None
    Path(dest).write_bytes(data)
    return None


def _write_pd_inputs(B, ws, facts, S) -> dict:
    """Scenarios and weights, rating buckets, the macro model and its scaling
    factors, and the GCC history and forecast behind the external scale."""
    chain = facts["chain"]
    out = {}
    B.title(ws, "PD inputs: scenarios, the macro model and the scaling factors",
            "From the run's frozen config (model.yml, model_inputs.yml) and "
            "static tables. Scenario weights are the resolved values the run "
            "priced with.")
    r = 4
    # rating -> bucket (Ratings.csv)
    ws.cell(r, 1, "Rating buckets (Ratings.csv)").font = B.f_sub
    B.header(ws, r + 1, 1, ["Rating type", "Rating", "Bucket (Hierarchy)", "Key"])
    rm = facts["ratings_map"]
    for k, m in enumerate(rm):
        rr = r + 2 + k
        B.put(ws, rr, 1, m["rt"], inp=True)
        B.put(ws, rr, 2, m["rating"], inp=True)
        B.put(ws, rr, 3, m["bucket"], "int", inp=True)
        B.put(ws, rr, 4, f"={_a(rr, 1, fixed=False)}&\"|\"&{_a(rr, 2, fixed=False)}")
    last = r + 1 + max(len(rm), 1)
    out["rm_key"] = _rng(r + 2, 4, last, 4, "PD inputs")
    out["rm_bucket"] = _rng(r + 2, 3, last, 3, "PD inputs")
    r = last + 2
    if chain is None:
        B.note(ws, r, 1, "The PD chain cannot be rebuilt from this run's config "
               f"({facts['chain_error']}); the ECL uses the run's StPD.csv.")
        out["chain"] = False
        return out
    out["chain"] = True
    scen = chain["scenarios"]
    nS = len(scen)
    # scenarios and weights
    ws.cell(r, 1, "Scenarios and weights").font = B.f_sub
    wext = chain["w_ext"]
    per_year = None
    if isinstance(wext, dict) and "per_year" in wext:
        per_year = wext["per_year"]
    ny = 0 if per_year is None else len(per_year)
    B.note(ws, r + 1, 1,
           f"Internal weights: mode '{chain['w_int_mode']}'. External weights: "
           f"mode '{chain['w_ext_mode']}', one set per forecast year; later "
           "years take the average.")
    hdr = ["Scenario", "Severity (z)", "Internal weight"] + \
        [f"External year {y}" for y in range(1, ny + 1)] + ["External average"]
    B.header(ws, r + 2, 1, hdr)
    out["scen_row"] = {}
    for k, (name, z) in enumerate(scen):
        rr = r + 3 + k
        out["scen_row"][name] = rr
        B.put(ws, rr, 1, name, inp=True)
        B.put(ws, rr, 2, z, "x", inp=True)
        B.put(ws, rr, 3, chain["w_int"].get(name), "pct", inp=True)
        for y in range(ny):
            B.put(ws, rr, 4 + y, float(per_year.iloc[y][name]), "pct", inp=True)
        if ny:
            B.put(ws, rr, 4 + ny, f"=AVERAGE({_rng(rr, 4, rr, 3 + ny)})", "pct")
        elif isinstance(wext, dict):
            B.put(ws, rr, 4, wext.get(name), "pct", inp=True)
    rr = r + 3 + nS
    B.put(ws, rr, 1, "Total", bold=True)
    for c in range(3, 5 + ny):
        B.put(ws, rr, c, f"=SUM({_rng(r + 3, c, r + 2 + nS, c)})", "pct")
    out["z"] = {name: _a(out["scen_row"][name], 2, "PD inputs") for name, _ in scen}
    out["w_int"] = {name: _a(out["scen_row"][name], 3, "PD inputs") for name, _ in scen}
    out["w_ext"] = {}
    for name, _ in scen:
        row = out["scen_row"][name]
        out["w_ext"][name] = ([_a(row, 4 + y, "PD inputs") for y in range(ny)],
                              _a(row, 4 + ny, "PD inputs"))
    r = rr + 2

    # internal macro model
    mevs = chain["mevs"]
    K = len(mevs)
    ws.cell(r, 1, f"Internal scale: macro model '{chain['model_id']}'").font = B.f_sub
    B.note(ws, r + 1, 1, "Per variable: stressed = forecast + z x sd x unit; "
           "logit = intercept + coefficient x stressed / unit; PD = e^logit / "
           "(e^logit + 1); SF = NORMSINV(PD) - NORMSINV(anchor). The scenario's "
           "SF is the weighted sum over the variables.")
    B.put(ws, r + 2, 1, "TTC anchor PD")
    B.put(ws, r + 2, 2, chain["anchor"], "pd", inp=True)
    out["anchor"] = _a(r + 2, 2, "PD inputs")
    B.put(ws, r + 3, 1, "Forecast years modelled")
    B.put(ws, r + 3, 2, chain["n_forecasts"], "int", inp=True)
    out["n_int"] = _a(r + 3, 2, "PD inputs")
    B.put(ws, r + 4, 1, "Last year of the curve (TTC reached)")
    B.put(ws, r + 4, 2, chain["max_maturity"], "int", inp=True)
    out["T"] = _a(r + 4, 2, "PD inputs")
    r += 6
    B.header(ws, r, 1, ["Variable", "Intercept", "Coefficient", "SD",
                        "Unit multiplier", "Weight"])
    mrow = []
    for k, m in enumerate(mevs):
        rr = r + 1 + k
        mrow.append(rr)
        B.put(ws, rr, 1, m["name"], inp=True)
        for c, key in enumerate(["intercept", "coefficient", "sd", "unit",
                                 "weight"], start=2):
            B.put(ws, rr, c, m[key], "rate", inp=True)
    r = r + 2 + K
    B.header(ws, r, 1, ["Forecast year"] + [m["name"] for m in mevs])
    frow = []
    fc = chain["forecasts"]
    for y in range(fc.shape[0]):
        rr = r + 1 + y
        frow.append(rr)
        B.put(ws, rr, 1, chain["forecast_years"][y], "int", inp=True)
        for k in range(K):
            B.put(ws, rr, 2 + k, float(fc[y, k]), "rate", inp=True)
    r = r + 2 + fc.shape[0]
    hdr = ["Scenario", "Year"]
    for m in mevs:
        nm = m["name"][:18]
        hdr += [f"{nm}: stressed", f"{nm}: logit", f"{nm}: PD", f"{nm}: SF"]
    hdr += ["Scenario SF"]
    B.header(ws, r, 1, hdr)
    out["sf_int"] = {}
    n = chain["n_forecasts"]
    rr = r + 1
    for name, _ in scen:
        for y in range(1, n + 1):
            B.put(ws, rr, 1, name)
            B.put(ws, rr, 2, y, "int")
            sfs = []
            for k in range(K):
                c0 = 3 + 4 * k
                mr = mrow[k]
                F = _a(frow[y - 1], 2 + k)
                B.put(ws, rr, c0,
                      f"={F}+{out['z'][name]}*{_a(mr, 4)}*{_a(mr, 5)}", "rate")
                B.put(ws, rr, c0 + 1,
                      f"={_a(mr, 2)}+{_a(mr, 3)}*{_a(rr, c0, fixed=False)}/{_a(mr, 5)}",
                      "rate")
                B.put(ws, rr, c0 + 2,
                      f"=EXP({_a(rr, c0 + 1, fixed=False)})/(EXP({_a(rr, c0 + 1, fixed=False)})+1)",
                      "pd")
                B.put(ws, rr, c0 + 3,
                      f"=_xlfn.NORM.S.INV({_a(rr, c0 + 2, fixed=False)})-"
                      f"_xlfn.NORM.S.INV({out['anchor']})", "rate")
                sfs.append(f"{_a(rr, c0 + 3, fixed=False)}*{_a(mr, 6)}")
            B.put(ws, rr, 3 + 4 * K, "=" + "+".join(sfs), "rate", key=True)
            out["sf_int"][(name, y)] = _a(rr, 3 + 4 * K, "PD inputs")
            rr += 1
    r = rr + 1

    # external scale
    if not chain["ext_on"]:
        B.note(ws, r, 1, "External scale: no GCC history or forecast in this "
               "run's config, so the external grades take the internal chain "
               "(as build_stpd_from_static does).")
        out["ext"] = False
        return out
    out["ext"] = True
    ws.cell(r, 1, "External scale: GCC real GDP growth").font = B.f_sub
    B.note(ws, r + 1, 1, "Each year's regional growth is the countries' growth "
           "weighted by that year's GDP at current prices. The scenario moves "
           "the forecast by z x the history's standard deviation; the SF is "
           "NORMSINV of where that lands in the history (PERCENTRANK.EXC, "
           "truncated to 3 significant digits; outside the history it takes "
           "1/(n+1) or n/(n+1)).")
    r += 2
    tabs = chain["gcc_tables"]
    hist = chain["gcc_hist"]
    h = tabs["hist"]
    if h is not None and len(h["years"]) == hist.size:
        cn = h["countries"]
        nc = len(cn)
        B.header(ws, r, 1, ["History year"] + [f"{c}: growth" for c in cn] +
                 [f"{c}: GDP" for c in cn] + ["GCC growth (GDP-weighted)"])
        for k, y in enumerate(h["years"]):
            rr = r + 1 + k
            B.put(ws, rr, 1, y, "int", inp=True)
            for j in range(nc):
                B.put(ws, rr, 2 + j, _num(h["growth"][k, j]), "rate", inp=True)
                B.put(ws, rr, 2 + nc + j, _num(h["gdp"][k, j]), "amt", inp=True)
            G = _rng(rr, 2, rr, 1 + nc)
            P = _rng(rr, 2 + nc, rr, 1 + 2 * nc)
            B.put(ws, rr, 2 + 2 * nc,
                  f"=SUMPRODUCT(--ISNUMBER({G}),--ISNUMBER({P}),{G},{P})/"
                  f"SUMPRODUCT(--ISNUMBER({G}),--ISNUMBER({P}),{P})", "rate")
        hcol = 2 + 2 * nc
    else:
        B.header(ws, r, 1, ["History (observation)", "GCC growth (GDP-weighted)"])
        for k, v in enumerate(hist):
            rr = r + 1 + k
            B.put(ws, rr, 1, k + 1, "int")
            B.put(ws, rr, 2, float(v), "rate", inp=True)
        hcol = 2
    H = _rng(r + 1, hcol, r + hist.size, hcol, "PD inputs")
    out["hist"] = H
    r = r + hist.size + 2
    B.put(ws, r, 1, "History: mean")
    B.put(ws, r, 2, f"=AVERAGE({H})", "rate")
    B.put(ws, r + 1, 1, "History: standard deviation")
    B.put(ws, r + 1, 2, f"=_xlfn.STDEV.S({H})", "rate")
    sd = _a(r + 1, 2, "PD inputs")
    B.put(ws, r + 2, 1, "History: observations")
    B.put(ws, r + 2, 2, f"=COUNT({H})", "int")
    nH = _a(r + 2, 2, "PD inputs")
    r += 4
    ext_fc = chain["ext_fc"]
    nE = int(ext_fc.size)
    out["n_ext"] = nE
    f = tabs["fc"]
    if f is not None and f["growth"].shape[0] == nE:
        cn = f["countries"]
        nc = len(cn)
        B.header(ws, r, 1, ["Forecast year"] + [f"{c}: growth" for c in cn] +
                 [f"{c}: GDP" for c in cn] + ["GCC growth (GDP-weighted)"])
        for k in range(nE):
            rr = r + 1 + k
            B.put(ws, rr, 1, k + 1, "int")
            for j in range(nc):
                B.put(ws, rr, 2 + j, _num(f["growth"][k, j]), "rate", inp=True)
                B.put(ws, rr, 2 + nc + j, _num(f["gdp"][k, j]), "amt", inp=True)
            Gr = _rng(rr, 2, rr, 1 + nc)
            P = _rng(rr, 2 + nc, rr, 1 + 2 * nc)
            B.put(ws, rr, 2 + 2 * nc,
                  f"=SUMPRODUCT(--ISNUMBER({Gr}),--ISNUMBER({P}),{Gr},{P})/"
                  f"SUMPRODUCT(--ISNUMBER({Gr}),--ISNUMBER({P}),{P})", "rate")
        fcol = 2 + 2 * nc
    else:
        B.header(ws, r, 1, ["Forecast year", "GCC growth"])
        for k in range(nE):
            rr = r + 1 + k
            B.put(ws, rr, 1, k + 1, "int")
            B.put(ws, rr, 2, float(ext_fc[k]), "rate", inp=True)
        fcol = 2
    fcrow = [r + 1 + k for k in range(nE)]
    r = r + nE + 2
    B.header(ws, r, 1, ["Scenario", "Year", "Stressed growth", "Rank (exclusive)",
                        "Rank, 3 significant digits", "Scenario SF"])
    out["sf_ext"] = {}
    rr = r + 1
    for name, _ in scen:
        for y in range(1, nE + 1):
            x = _a(rr, 3, fixed=False)
            B.put(ws, rr, 1, name)
            B.put(ws, rr, 2, y, "int")
            B.put(ws, rr, 3, f"={_a(fcrow[y - 1], fcol)}+{out['z'][name]}*{sd}", "rate")
            B.put(ws, rr, 4, f"=IF({x}<MIN({H}),1/({nH}+1),IF({x}>MAX({H}),"
                             f"{nH}/({nH}+1),_xlfn.PERCENTRANK.EXC({H},{x},15)))", "rate")
            p = _a(rr, 4, fixed=False)
            B.put(ws, rr, 5, f"=ROUNDDOWN({p},2-INT(LOG10({p})))", "rate")
            B.put(ws, rr, 6, f"=_xlfn.NORM.S.INV({_a(rr, 5, fixed=False)})", "rate",
                  key=True)
            out["sf_ext"][(name, y)] = _a(rr, 6, "PD inputs")
            rr += 1
    B.widths(ws, {"A": 26, "B": 14, "C": 16, "D": 16, "E": 16, "F": 16})
    return out


def _write_pd_annual(B, ws, facts, pin) -> dict:
    """Per chain grade: PiT, cumulative and marginal PD per scenario and year,
    and the scenario-weighted marginal."""
    chain = facts["chain"]
    B.title(ws, "PD annual: from the scaling factors to the weighted marginal PD",
            "Years up to the forecast horizon: internal PiT = NORMSDIST(NORMSINV"
            "(TTC) + SF); external (Basel ASRF) PiT = NORMSDIST((NORMSINV(TTC) - "
            "SQRT(R) x SF) / SQRT(1 - R)), R = 0.24 - 0.12 (1 - e^(-50 TTC)) / "
            "(1 - e^(-50)). Later years revert to TTC: LN(f x (e^TTC - "
            "e^PiT_n) + e^PiT_n), f = (t - n) / (T - n). Cumulative = 1 - (1 - "
            "previous) x (1 - PiT); marginal = cumulative - previous. The "
            "weighted marginal sums the scenarios' marginals x their weights.")
    out = {"blocks": {}}
    if not pin.get("chain"):
        B.note(ws, 4, 1, "Not available for this run: see PD inputs.")
        return out
    scen = chain["scenarios"]
    T = chain["max_maturity"]
    r = 4
    done = {}
    for c in facts["curves"]:
        if c["chain"] is None:
            continue
        ext = c["rt"] == 2 and pin.get("ext")
        bkey = (c["rt"], c["chain"][0])
        if bkey in done:
            out["blocks"][c["key"]] = done[bkey]
            continue
        rating, ttc = c["chain"]
        scale = "External" if c["rt"] == 2 else "Internal"
        ws.cell(r, 1, f"{scale} grade {rating}").font = B.f_sub
        B.put(ws, r + 1, 1, "TTC PD")
        B.put(ws, r + 1, 2, ttc, "pd", inp=True)
        TTC = _a(r + 1, 2)
        n = pin["n_ext"] if ext else chain["n_forecasts"]
        B.put(ws, r + 2, 1, "Forecast years (n)")
        B.put(ws, r + 2, 2, n, "int")
        N = _a(r + 2, 2)
        B.put(ws, r + 3, 1, "Last year (T)")
        B.put(ws, r + 3, 2, f"={pin['T']}", "int")
        TT = _a(r + 3, 2)
        if ext:
            B.put(ws, r + 4, 1, "Asset correlation R")
            B.put(ws, r + 4, 2, f"=0.24-0.12*(1-EXP(-50*{TTC}))/(1-EXP(-50))", "rate")
            RR = _a(r + 4, 2)
        hr = r + 6
        hdr = ["Year"]
        for name, _ in scen:
            hdr += [f"{name}: SF", f"{name}: PiT", f"{name}: cumulative",
                    f"{name}: marginal", f"{name}: weight"]
        hdr += ["Weighted marginal"]
        B.header(ws, hr, 1, hdr)
        first = hr + 1
        for y in range(1, T + 1):
            rr = hr + y
            B.put(ws, rr, 1, y, "int")
            terms = []
            for k, (name, _) in enumerate(scen):
                c0 = 2 + 5 * k
                sfc, pit, cum, mar, wc = (_a(rr, c0 + i, fixed=False) for i in range(5))
                if y <= n:
                    sf = pin["sf_ext"][(name, y)] if ext else pin["sf_int"][(name, y)]
                    B.put(ws, rr, c0, f"={sf}", "rate")
                    if ext:
                        f = (f"=IF({TTC}<=0,0,IF({TTC}>=1,1,_xlfn.NORM.S.DIST(("
                             f"_xlfn.NORM.S.INV({TTC})-SQRT({RR})*{sfc})/SQRT(1-{RR}),TRUE)))")
                    else:
                        f = (f"=IF({TTC}=0,0,_xlfn.NORM.S.DIST(_xlfn.NORM.S.INV({TTC})"
                             f"+{sfc},TRUE))")
                else:
                    anchor = _a(hr + n, c0 + 1)
                    f = (f"=IF({TTC}=0,0,LN(({_a(rr, 1, fixed=False)}-{N})/({TT}-{N})*"
                         f"(EXP({TTC})-EXP({anchor}))+EXP({anchor})))")
                B.put(ws, rr, c0 + 1, f, "pd")
                if y == 1:
                    B.put(ws, rr, c0 + 2, f"={pit}", "pd")
                    B.put(ws, rr, c0 + 3, f"={cum}", "pd")
                else:
                    prev = _a(rr - 1, c0 + 2, fixed=False)
                    B.put(ws, rr, c0 + 2, f"=1-(1-{prev})*(1-{pit})", "pd")
                    B.put(ws, rr, c0 + 3, f"={cum}-{prev}", "pd")
                if c["rt"] == 2:            # the external scale's own weights
                    years, avg = pin["w_ext"][name]
                    wref = years[y - 1] if y <= len(years) else avg
                else:
                    wref = pin["w_int"][name]
                B.put(ws, rr, c0 + 4, f"={wref}", "pct")
                terms.append(f"{mar}*{wc}")
            B.put(ws, rr, 2 + 5 * len(scen), "=" + "+".join(terms), "pd", key=True)
        blk = {"wm": (first, 2 + 5 * len(scen), first + T - 1), "T": T}
        out["blocks"][c["key"]] = blk
        done[bkey] = blk
        r = hr + T + 3
    B.widths(ws, {"A": 22, "B": 14})
    ws.freeze_panes = ws.cell(5, 2)
    return out


def _write_pd_monthly(B, ws, facts, pa, S) -> dict:
    """Months 0-600 per curve: the curve the ECL uses, the rebuilt curve, the
    run's StPD.csv and the difference."""
    curves = facts["curves"]
    nC = len(curves)
    B.title(ws, "PD monthly: cumulative PD by month",
            "Rebuilt: each year's weighted marginal is spread evenly over its "
            "12 months and accumulated (capped at the PD cap); after the last "
            "year the last year's marginal repeats. 'Used' is the curve the "
            "ECL tab reads: the rebuilt one or the run's StPD.csv, per the "
            "Settings tab.")
    hr = 5
    B.put(ws, 4, 1, "Curve (portfolio|bucket)", bold=True)
    labels = ["Month"]
    for blk in ("Used", "Rebuilt", "Run StPD", "Rebuilt - StPD"):
        labels += [f"{blk}: {c['label']}" for c in curves]
    B.header(ws, hr, 1, labels)
    # the keys the Lines tab matches on, one per curve, above the Used block
    for k, c in enumerate(curves):
        B.put(ws, 4, 2 + k, c["label"], inp=True)
    keys = _rng(4, 2, 4, 1 + max(nC, 1), "PD monthly")
    u0, r0, s0, d0 = 2, 2 + nC, 2 + 2 * nC, 2 + 3 * nC
    first = hr + 1
    for m in range(0, MAX_MONTH + 1):
        rr = first + m
        B.put(ws, rr, 1, m, "int")
        for k, c in enumerate(curves):
            blk = pa["blocks"].get(c["key"])
            if blk is not None:
                if m == 0:
                    f = "=0"
                else:
                    a, col, b = blk["wm"]
                    idx = f"MIN(ROUNDUP({_a(rr, 1, fixed=False)}/12,0),{blk['T']})"
                    f = (f"=MIN({S['pdcap']},{_a(rr - 1, r0 + k, fixed=False)}+"
                         f"INDEX({_rng(a, col, b, col, 'PD annual')},{idx})/12)")
                B.put(ws, rr, r0 + k, f, "pd")
            sp = c["stpd"]
            v = float(sp[m]) if sp is not None and m < len(sp) else None
            B.put(ws, rr, s0 + k, v, "pd", inp=True)
            rb, sv = _a(rr, r0 + k, fixed=False), _a(rr, s0 + k, fixed=False)
            B.put(ws, rr, d0 + k,
                  f"=IF(OR({rb}=\"\",{sv}=\"\"),\"\",{rb}-{sv})", "pd")
            B.put(ws, rr, u0 + k,
                  f"=IF({S['pdsrc']}=\"{PD_SOURCES[0]}\",{rb},{sv})", "pd")
    last = first + MAX_MONTH
    ws.freeze_panes = ws.cell(first, 2)
    B.widths(ws, {"A": 8})
    return {"keys": keys, "used": _rng(first, u0, last, u0 + max(nC, 1) - 1,
                                       "PD monthly"),
            "diff": _rng(first, d0, last, d0 + max(nC, 1) - 1, "PD monthly"),
            "first": first}


def _write_lgd(B, ws, facts, L, ln, S) -> dict:
    co = facts["collateral"]
    a, c, t = co["alloc"], co["coll"], co["types"]
    B.title(ws, "LGD: collateral to loss given default",
            "Net collateral per allocation = value x allocation x (1 - haircut) "
            "(nil when any of the three is missing). LGD = base x MAX(floor, "
            "(exposure - net collateral) / exposure); the floor keeps at least "
            "that share unsecured. Zero or missing exposure takes the setting's "
            "LGD.")
    r = 4
    ws.cell(r, 1, "Collateral (Collateral.csv)").font = B.f_sub
    B.header(ws, r + 1, 1, ["Collateral id", "Type", "Value"])
    cr0 = r + 2
    for k, row in enumerate(c.itertuples(index=False)):
        B.put(ws, cr0 + k, 1, row.collateral, inp=True)
        B.put(ws, cr0 + k, 2, row.type, inp=True)
        B.put(ws, cr0 + k, 3, _num(row.value), "amt", inp=True)
    cr1 = cr0 + max(len(c), 1) - 1
    ws.cell(r, 5, "Collateral types (CollateralType.csv)").font = B.f_sub
    B.header(ws, r + 1, 5, ["Type", "Name", "Haircut"])
    for k, row in enumerate(t.itertuples(index=False)):
        B.put(ws, cr0 + k, 5, row.type, inp=True)
        B.put(ws, cr0 + k, 6, _str(row.name), inp=True)
        B.put(ws, cr0 + k, 7, _num(row.haircut), "pct", inp=True)
    tr1 = cr0 + max(len(t), 1) - 1
    r = max(cr1, tr1) + 3
    ws.cell(r, 1, "Allocations (AccountCollateralAllocation.csv)").font = B.f_sub
    B.header(ws, r + 1, 1, ["Contract", "Collateral id", "Allocation",
                            "Type", "Value", "Haircut", "Net collateral"])
    ar0 = r + 2
    for k, row in enumerate(a.itertuples(index=False)):
        rr = ar0 + k
        B.put(ws, rr, 1, row.contract, inp=True)
        B.put(ws, rr, 2, row.collateral, inp=True)
        B.put(ws, rr, 3, _num(row.share), "pct", inp=True)
        idr = _a(rr, 2, fixed=False)
        B.put(ws, rr, 4, f"=IFERROR(INDEX({_rng(cr0, 2, cr1, 2)},MATCH({idr},"
                         f"{_rng(cr0, 1, cr1, 1)},0)),\"\")")
        B.put(ws, rr, 5, f"=IFERROR(INDEX({_rng(cr0, 3, cr1, 3)},MATCH({idr},"
                         f"{_rng(cr0, 1, cr1, 1)},0)),\"\")", "amt")
        B.put(ws, rr, 6, f"=IFERROR(INDEX({_rng(cr0, 7, tr1, 7)},MATCH("
                         f"{_a(rr, 4, fixed=False)},{_rng(cr0, 5, tr1, 5)},0)),\"\")",
              "pct")
        v, s_, h = (_a(rr, i, fixed=False) for i in (5, 3, 6))
        B.put(ws, rr, 7, f"=IF(COUNT({v},{s_},{h})=3,{v}*{s_}*(1-{h}),0)", "amt")
    ar1 = ar0 + max(len(a), 1) - 1
    r = ar1 + 3
    ws.cell(r, 1, "LGD per position").font = B.f_sub
    B.header(ws, r + 1, 1, ["Line", "Contract", "On balance", "Net collateral",
                            "Collateral cover", "LGD", "Report LGD (Stage 1/2)"])
    out = {}
    for k, l in enumerate(L):
        rr = r + 2 + k
        x = l["line"]
        B.put(ws, rr, 1, f"={ln(x, 'Line')}")
        B.put(ws, rr, 2, f"={ln(x, 'Contract')}")
        onb = ln(x, "On balance")
        B.put(ws, rr, 3, f"=IF(ISNUMBER({onb}),{onb},0)", "amt")
        B.put(ws, rr, 4, f"=SUMIFS({_rng(ar0, 7, ar1, 7)},{_rng(ar0, 1, ar1, 1)},"
                         f"{_a(rr, 2, fixed=False)})", "amt")
        e, n = _a(rr, 3, fixed=False), _a(rr, 4, fixed=False)
        B.put(ws, rr, 5, f"=IF({e}>0,{n}/{e},\"\")", "pct")
        B.put(ws, rr, 6, f"=IF({e}>0,{S['lgd_base']}*MAX({S['lgd_floor']},"
                         f"({e}-{n})/{e}),{S['zero_lgd']})", "pct", key=True)
        B.put(ws, rr, 7, l["rep_lgd"], "pct", inp=True)
        out[x] = {"lgd": _a(rr, 6, "LGD"), "net": _a(rr, 4, "LGD"),
                  "onb": _a(rr, 3, "LGD"), "cover": _a(rr, 5, "LGD")}
    B.widths(ws, {"A": 16, "B": 16, "C": 16, "D": 16, "E": 14, "F": 14, "G": 14})
    return out


def _write_ead(B, ws, facts, L, ln, stage, S) -> dict:
    B.title(ws, "EAD: the exposure path and the horizon",
            "A position with a repayment schedule uses its monthly EAD from "
            "LifeTimeParameterOther.csv (ECL tab, blue). Without one: months to "
            "maturity N (a part month counts; 3 at or past maturity), the "
            "shape from the Settings rules, payments every f months counted "
            "back from maturity after any deferral d, so n = ROUNDUP((N - d) / "
            "f) payments with the first in month p1 = N - (n - 1) f. Linear: "
            "EAD = balance x (1 - paid / n). Annuity: balance x ((1 + r f)^n - "
            "(1 + r f)^paid) / ((1 + r f)^n - 1), r = NIR / 12. Bullet: the "
            "balance. Horizon: 12 months at Stage 1 (or less), the whole path "
            "otherwise, at most month 600.")
    cols = ["Line", "Contract", "Source", "Months to maturity (N)", "Shape",
            "f", "d", "Amortising months", "n payments", "First payment p1",
            "Monthly rate r", "Path length", "Stage", "Horizon (months)"]
    E = {n: i + 1 for i, n in enumerate(cols)}
    B.header(ws, 4, 1, cols)
    out = {}
    for k, l in enumerate(L):
        rr = 5 + k
        x = l["line"]
        g = lambda name: _a(rr, E[name], fixed=False)
        B.put(ws, rr, E["Line"], f"={ln(x, 'Line')}")
        B.put(ws, rr, E["Contract"], f"={ln(x, 'Contract')}")
        sched = l["schedule"]
        B.put(ws, rr, E["Source"], "Schedule" if sched is not None else "Fallback")
        mat, ext = ln(x, "Maturity date"), ln(x, "Extract date")
        mm = (f"((YEAR({mat})-YEAR({ext}))*12+MONTH({mat})-MONTH({ext})+"
              f"IF(DAY({mat})>DAY({ext}),1,0))")
        B.put(ws, rr, E["Months to maturity (N)"],
              f"=IF(OR({mat}=\"\",{ext}=\"\"),{S['minm']},IF({mm}>0,{mm},{S['minm']}))",
              "int")
        N = g("Months to maturity (N)")
        pf, pt = ln(x, "Portfolio"), ln(x, "Payment type")
        shape = (f"IFERROR(INDEX({S['rule_shape']},MATCH({pf}&\"|\"&{pt},"
                 f"{S['rule_key']},0)),IFERROR(INDEX({S['rule_shape']},MATCH("
                 f"\"|\"&{pt},{S['rule_key']},0)),{S['default_shape']}))")
        B.put(ws, rr, E["Shape"],
              f"=IF({g('Source')}=\"Schedule\",\"\",IF({N}<={S['minm']},\"bullet\",{shape}))")
        pfq, dfr = ln(x, "Payment frequency (months)"), ln(x, "Deferral (months)")
        B.put(ws, rr, E["f"], f"=MAX(1,IF(ISNUMBER({pfq}),TRUNC({pfq}),1))", "int")
        B.put(ws, rr, E["d"], f"=MAX(0,MIN(IF(ISNUMBER({dfr}),TRUNC({dfr}),0),"
                              f"MAX(1,{N})-1))", "int")
        B.put(ws, rr, E["Amortising months"], f"=MAX(1,MAX(1,{N})-{g('d')})", "int")
        B.put(ws, rr, E["n payments"],
              f"=MAX(1,ROUNDUP({g('Amortising months')}/{g('f')},0))", "int")
        B.put(ws, rr, E["First payment p1"],
              f"=MAX(1,{N})-({g('n payments')}-1)*{g('f')}", "int")
        nir = ln(x, "Nominal interest rate")
        B.put(ws, rr, E["Monthly rate r"], f"=IF(ISNUMBER({nir}),{nir}/12,0)", "rate")
        if sched is not None:
            B.put(ws, rr, E["Path length"], len(sched), "int")
        else:
            B.put(ws, rr, E["Path length"], f"=MIN(MAX(1,{N}),{S['maxm']})", "int")
        B.put(ws, rr, E["Stage"], f"={stage(x)}", "int")
        B.put(ws, rr, E["Horizon (months)"],
              f"=MIN(IF({g('Stage')}=1,MIN({S['s1m']},{g('Path length')}),"
              f"{g('Path length')}),{S['maxm']})", "int", key=True)
        out[x] = {c: _a(rr, E[c], "EAD") for c in cols}
    B.widths(ws, {"A": 6, "B": 14, "C": 10, "D": 11, "E": 10, "F": 5, "G": 5,
                  "H": 11, "I": 10, "J": 11, "K": 12, "L": 10, "M": 7, "N": 10})
    ws.freeze_panes = ws.cell(5, 3)
    return out


def _write_ecl(B, ws, facts, L, ln, lg, ead, pm, S) -> dict:
    B.title(ws, "ECL month by month",
            "ECL(t) = EAD(t) x LGD x (cumPD(t+1) - cumPD(t)) / (1 + EIR)^(t/12) "
            "for t below the horizon; the position's ECL is the sum (Summary "
            "tab), capped at the exposure, with Stage 3 per the Settings.")
    cols = ["Line", "Contract", "t (month)", "EAD(t)", "In horizon",
            "cumPD(t)", "cumPD(t+1)", "Marginal PD", "LGD", "EIR",
            "Discount factor", "ECL(t)"]
    B.header(ws, 4, 1, cols)
    rr = 5
    out = {}
    for l in L:
        x = l["line"]
        e = ead[x]
        sched = l["schedule"]
        if sched is not None:
            n = min(len(sched), MAX_MONTH)
        else:
            n = min(max(1, _fallback_n(l)), MAX_MONTH)
        start = rr
        onb = ln(x, "On balance")
        curve = ln(x, "Curve no.")
        for t in range(n):
            B.put(ws, rr, 1, x, "int")
            B.put(ws, rr, 2, l["cid"])
            B.put(ws, rr, 3, t, "int")
            T = _a(rr, 3, fixed=False)
            if sched is not None:
                B.put(ws, rr, 4, float(sched[t]), "amt", inp=True)
            else:
                paid = (f"IF({T}<{e['First payment p1']},0,MIN({e['n payments']},"
                        f"INT(({T}-{e['First payment p1']})/{e['f']})+1))")
                rp = f"({e['Monthly rate r']}*{e['f']})"
                coef = (f"IF({e['Shape']}=\"linear\",MAX(0,1-{paid}/{e['n payments']}),"
                        f"IF(AND({e['Shape']}=\"annuity\",{e['Monthly rate r']}>1E-9),"
                        f"IF({paid}>={e['n payments']},0,MAX(0,((1+{rp})^{e['n payments']}"
                        f"-(1+{rp})^{paid})/((1+{rp})^{e['n payments']}-1))),"
                        f"IF({e['Shape']}=\"annuity\",MAX(0,1-{paid}/{e['n payments']}),1)))")
                B.put(ws, rr, 4, f"=IF(ISNUMBER({onb}),{onb},0)*{coef}", "amt")
            B.put(ws, rr, 5, f"={T}<{e['Horizon (months)']}")
            B.put(ws, rr, 6, f"=INDEX({pm['used']},{T}+1,{curve})", "pd")
            B.put(ws, rr, 7, f"=INDEX({pm['used']},{T}+2,{curve})", "pd")
            B.put(ws, rr, 8, f"={_a(rr, 7, fixed=False)}-{_a(rr, 6, fixed=False)}", "pd")
            B.put(ws, rr, 9, f"={lg[x]['lgd']}", "pct")
            eir = ln(x, "EIR")
            B.put(ws, rr, 10, f"={eir}", "rate")
            B.put(ws, rr, 11, f"=IF(ISNUMBER({eir}),(1+{eir})^({T}/12),NA())", "rate")
            B.put(ws, rr, 12,
                  f"=IF({_a(rr, 5, fixed=False)},{_a(rr, 4, fixed=False)}*"
                  f"{_a(rr, 9, fixed=False)}*{_a(rr, 8, fixed=False)}/"
                  f"{_a(rr, 11, fixed=False)},0)", "amt")
            rr += 1
        out[x] = (start, rr - 1)
    out["_col_line"] = _rng(5, 1, max(rr - 1, 5), 1, "ECL")
    out["_col_ecl"] = _rng(5, 12, max(rr - 1, 5), 12, "ECL")
    ws.auto_filter.ref = f"A4:L{max(rr - 1, 5)}"
    ws.freeze_panes = ws.cell(5, 4)
    B.widths(ws, {"A": 6, "B": 14, "C": 9, "D": 16, "E": 10, "F": 13, "G": 13,
                  "H": 13, "I": 10, "J": 10, "K": 12, "L": 14})
    return out


def _fallback_n(l) -> int:
    from .engine import months_to_maturity
    return months_to_maturity(l["d_maturity"], l["d_extract"])


def _write_summary(B, ws, facts, L, ln, stage, lg, ead, ecl, pm, S, GR, G):
    B.title(ws, f"ECL calculator: customer {facts['customer_id']}"
                + (f" ({facts['customer_name']})" if facts["customer_name"] else ""))
    info = [
        ("Run", facts["run_id"]),
        ("Extract date", L[0]["d_extract"] if L else None),
        ("Engine", facts["engine"]),
        ("Report compared", facts["report_file"] or "none in the run"),
        ("Overlay report", facts["overlay_file"] or "none"),
        ("PD curve used", f"={S['pdsrc']}"),
        ("Rebuilt vs run StPD, largest gap", f"=IFERROR(MAX(ABS(MIN({pm['diff']})),"
                                             f"ABS(MAX({pm['diff']}))),\"\")"),
        ("Generated", facts["generated_at"]),
    ]
    for k, (lab, v) in enumerate(info):        # label overflows into B:C
        B.put(ws, 3 + k, 1, lab, bold=True)
        c = B.put(ws, 3 + k, 4, v, "date" if lab == "Extract date" else
                  ("pd" if lab.startswith("Rebuilt") else None))
        c.alignment = B.Align(horizontal="left")
    B.note(ws, 12, 1, "How to read it: Lines (inputs) -> Overrides -> Staging -> "
           "LGD -> PD inputs -> PD annual -> PD monthly -> EAD -> ECL. Blue "
           "cells are the run's inputs; black cells are formulas. Change an "
           "input and the provision follows; the report column stays as booked.")
    ws.cell(21, 1, "Per position").font = B.f_sub
    hr = 22
    cols = ["Line", "Book", "Contract", "Portfolio", "Rating", "Stage", "Why",
            "On balance", "Net collateral", "LGD", "PD lifetime", "Horizon",
            "ECL (this workbook)", "ECL (report)", "Difference", "Overlay",
            "ECL final", "Report final", "Check"]
    U = {n: i + 1 for i, n in enumerate(cols)}
    B.header(ws, hr, 1, cols)
    first = hr + 1
    for k, l in enumerate(L):
        rr = first + k
        x = l["line"]
        g = lambda name: _a(rr, U[name], fixed=False)
        B.put(ws, rr, U["Line"], f"={ln(x, 'Line')}")
        B.put(ws, rr, U["Book"], f"={ln(x, 'Book')}")
        B.put(ws, rr, U["Contract"], f"={ln(x, 'Contract')}")
        B.put(ws, rr, U["Portfolio"], f"={ln(x, 'Portfolio')}")
        B.put(ws, rr, U["Rating"], f"={ln(x, 'Rating')}")
        B.put(ws, rr, U["Stage"], f"={stage(x)}", "int")
        B.put(ws, rr, U["Why"], f"={_a(GR[x], G['Why'], 'Staging')}")
        B.put(ws, rr, U["On balance"], f"={lg[x]['onb']}", "amt")
        B.put(ws, rr, U["Net collateral"], f"={lg[x]['net']}", "amt")
        B.put(ws, rr, U["LGD"], f"={lg[x]['lgd']}", "pct")
        e = ead[x]
        h = (f"MIN(MAX(0,IF({g('Stage')}=1,MIN({S['s1m']},{e['Months to maturity (N)']}),"
             f"{e['Months to maturity (N)']})),{S['maxm']})")
        B.put(ws, rr, U["PD lifetime"],
              f"=IFERROR(INDEX({pm['used']},{h}+1,{ln(x, 'Curve no.')}),\"\")", "pd")
        B.put(ws, rr, U["Horizon"], f"={e['Horizon (months)']}", "int")
        a, b = ecl[x]
        tot = f"SUM({_rng(a, 12, b, 12, 'ECL')})"
        onb = g("On balance")
        B.put(ws, rr, U["ECL (this workbook)"],
              f"=IF({g('Stage')}=3,IF({S['stage3']}=\"zero\",0,MAX(0,{onb})),"
              f"IF(AND({S['cap']},{onb}>0),MIN({tot},{onb}),{tot}))", "amt", key=True)
        B.put(ws, rr, U["ECL (report)"], l["rep_ecl"], "amt", inp=True)
        B.put(ws, rr, U["Difference"],
              f"=IF({g('ECL (report)')}=\"\",\"\",{g('ECL (this workbook)')}-"
              f"{g('ECL (report)')})", "amt")
        B.put(ws, rr, U["Overlay"], l["overlay"] if l["overlay"] is not None else 0,
              "amt", inp=True)
        B.put(ws, rr, U["ECL final"], f"={g('ECL (this workbook)')}+{g('Overlay')}",
              "amt")
        B.put(ws, rr, U["Report final"], l["rep_final"], "amt", inp=True)
        tol = lambda a, b: f"ABS({a}-{b})<=MAX(0.01,1E-9*ABS({b}))"
        B.put(ws, rr, U["Check"],
              f"=IF({g('ECL (report)')}=\"\",\"not in report\",IF(AND("
              f"{tol(g('ECL (this workbook)'), g('ECL (report)'))},OR("
              f"{g('Report final')}=\"\",{tol(g('ECL final'), g('Report final'))})),"
              f"\"OK\",\"DIFFERS\"))")
    last = first + max(len(L), 1) - 1
    # totals by stage
    ws.cell(14, 1, "Totals").font = B.f_sub
    B.header(ws, 15, 1, ["Stage", "Positions", "On balance",
                         "ECL (this workbook)", "ECL (report)", "Difference",
                         "ECL final", "Report final"])
    col = lambda n: _rng(first, U[n], last, U[n])
    for k, s in enumerate([1, 2, 3, "Total"]):
        rr = 16 + k
        B.put(ws, rr, 1, s, bold=s == "Total")
        if s == "Total":
            f = lambda n: f"=SUM({_rng(16, n, 18, n)})"
            for c in range(2, 9):
                B.put(ws, rr, c, f(c), "int" if c == 2 else "amt", bold=True)
            continue
        B.put(ws, rr, 2, f"=COUNTIF({col('Stage')},{s})", "int")
        for c, n in ((3, "On balance"), (4, "ECL (this workbook)"),
                     (5, "ECL (report)"), (7, "ECL final"), (8, "Report final")):
            B.put(ws, rr, c, f"=SUMIFS({col(n)},{col('Stage')},{s})", "amt")
        B.put(ws, rr, 6, f"={_a(rr, 4, fixed=False)}-{_a(rr, 5, fixed=False)}", "amt")
    from openpyxl.formatting.rule import CellIsRule
    rng = f"{_col_letter(U['Check'])}{first}:{_col_letter(U['Check'])}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=['"OK"'],
        fill=B.Fill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=['"DIFFERS"'],
        fill=B.Fill("solid", fgColor="FFC7CE")))
    B.widths(ws, {"A": 8, "B": 12, "C": 16, "D": 18, "E": 10, "F": 7, "G": 34,
                  "H": 16, "I": 16, "J": 9, "K": 12, "L": 9, "M": 18, "N": 18,
                  "O": 12, "P": 12, "Q": 16, "R": 16, "S": 10})
    ws.freeze_panes = ws.cell(first, 4)


def customer_ecl_workbook(run_dir, customer_id, dest=None):
    """The calculator for one customer of a run: bytes, or written to
    ``dest``."""
    return write_customer_workbook(customer_ecl_facts(run_dir, customer_id), dest)
